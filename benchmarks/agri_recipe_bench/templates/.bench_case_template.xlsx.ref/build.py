from __future__ import annotations

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

OUT = Path(r"C:\Users\xing\Desktop\智慧农业生态\benchmarks\agri_recipe_bench\templates\bench_case_template.xlsx")
OUT.parent.mkdir(parents=True, exist_ok=True)

wb = Workbook()
wb.properties.title = "AgriRecipeBench Eval-First Gate Template"

font_title = Font(bold=True, size=14, color="FF1F2937")
font_header = Font(bold=True, color="FFFFFFFF")
font_section = Font(bold=True, color="FF111827")
font_note = Font(italic=True, color="FF6B7280")
font_mono = Font(name="Consolas", size=10, color="FF111827")
font_body = Font(color="FF111827")

fill_title = PatternFill("solid", fgColor="FFEFF6FF")
fill_header = PatternFill("solid", fgColor="FF3B82F6")
fill_section = PatternFill("solid", fgColor="FFF3F4F6")
fill_input = PatternFill("solid", fgColor="FFECFDF5")
fill_check = PatternFill("solid", fgColor="FFFEF3C7")
fill_warn = PatternFill("solid", fgColor="FFFEE2E2")

border_side = Side(style="thin", color="FFD1D5DB")
border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="top", wrap_text=True)


def style_range(ws, start_row, end_row, start_col, end_col, *, fill=None, font=None, align=None):
    for row in ws.iter_rows(min_row=start_row, max_row=end_row, min_col=start_col, max_col=end_col):
        for cell in row:
            cell.border = border
            cell.alignment = align or left
            if fill:
                cell.fill = fill
            if font:
                cell.font = font


def set_widths(ws, widths):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

# ---------- Sheet 1: README ----------
ws = wb.active
ws.title = "README"
ws.sheet_view.showGridLines = False
ws["A1"] = "AgriRecipeBench Eval-First Gate Template"
ws["A1"].font = font_title
ws.merge_cells("A1:H1")
ws["A2"] = "用途：新增科学模块进入主干前，先用本表登记 4 轴评测用例，再提交代码。"
ws["A2"].alignment = left
ws.merge_cells("A2:H2")

ws["A4"] = "适用对象"
ws["B4"] = "新增科学模块：气候引擎、作物模型、无土基质模型、营养/病虫害/物候相关新内核"
ws["A4"].font = font_section
ws["B4"].alignment = left
ws.merge_cells("B4:H4")
ws["A5"] = "最低要求"
ws["B5"] = "每条新模块至少补齐 target_prediction / state_rolling / counterfactual / generalization 四轴各 1 条用例"
ws["A5"].font = font_section
ws["B5"].alignment = left
ws.merge_cells("B5:H5")
ws["A6"] = "提交前"
ws["B6"] = "运行 `python scripts/check_bench_coverage.py` 与 `python benchmarks/agri_recipe_bench/run_bench.py`"
ws["A6"].font = font_section
ws["B6"].alignment = left
ws.merge_cells("B6:H6")

ws["A8"] = "使用说明"
ws["A8"].font = font_section
style_range(ws, 8, 8, 1, 8, fill=fill_section, align=center)
ws["A9"] = "1"
ws["B9"] = "把模板文件复制到 benchmarks/agri_recipe_bench/cases/ 下的 JSON 用例文件。"
ws["C9"] = "文件名建议包含模块名与轴名，例如 case_substrate_target.json。"
ws["A10"] = "2"
ws["B10"] = "按当前用例轴填写 id / axis / input / expected / perturbation / provenance。"
ws["C10"] = "只填本轴必需的字段；不要把所有轴一次性塞进一个 JSON。"
ws["A11"] = "3"
ws["B11"] = "用现有 runner 直接消费；不要改 runner 去迁就临时字段。"
ws["C11"] = "如果新模块新增输出字段，先补 runner，再补用例。"

ws["A13"] = "备注"
ws["A13"].font = font_note
ws["A13"].alignment = left
ws["A13"] = "模板本身不进入 cases/，不会被 runner 误扫。"
ws["A14"] = "若用例需要真实 env_recipe，请在 input 中同时给出 env_recipe_path 与最小可读的 env_recipe 片段。"

set_widths(ws, {"A": 12, "B": 42, "C": 28, "D": 20, "E": 20, "F": 20, "G": 16, "H": 16})
ws.row_dimensions[1].height = 26
ws.row_dimensions[8].height = 22
ws.row_dimensions[9].height = 34
ws.row_dimensions[10].height = 34
ws.row_dimensions[11].height = 34

# ---------- Sheet 2: 字段字典 ----------
ws2 = wb.create_sheet("字段字典")
ws2.sheet_view.showGridLines = False
headers = ["字段", "类型", "必填", "适用轴", "示例", "说明"]
for i, h in enumerate(headers, start=1):
    ws2.cell(row=1, column=i, value=h)
style_range(ws2, 1, 1, 1, len(headers), fill=fill_header, font=font_header, align=center)

rows = [
    ("id", "文本", "是", "全部", "case_substrate_target_1", "唯一标识；建议 case_<module>_<axis>_<short_name>"),
    ("axis", "枚举", "是", "全部", "target_prediction", "只能选 target_prediction / state_rolling / counterfactual / generalization"),
    ("description", "文本", "是", "全部", "验证基质对产量的影响", "说明用例验证什么"),
    ("input", "对象", "是", "全部", "{}", "基准输入；建议保留 crop / zone_id / start_date / adapt_score / growth_plan / env_recipe 或 env_recipe_path"),
    ("perturbation", "对象", "否", "counterfactual", '{"adapt_score":{"before":0.9,"after":0.3}}', "只在反事实轴使用；给出 before / after"),
    ("expected", "对象", "是", "全部", "{}", "按当前轴填写判据；不要混写其他轴的 expected"),
    ("provenance", "对象", "是", "全部", '{"note":"为什么需要这条用例"}', "记录判定动机，便于回看"),
    ("crop", "文本", "否", "target_prediction", "番茄", "作物名，若为跨轴模板可省略"),
    ("zone_id", "文本", "否", "target_prediction", "mediterranean", "单区用例时填写；跨分区用例请用 zones 数组"),
    ("start_date", "日期", "否", "target_prediction", "2026-03-01", "建议 ISO 日期字符串"),
    ("adapt_score", "数字", "否", "target_prediction", "0.85", "适应度输入"),
    ("env_recipe_path", "文本", "否", "target_prediction", "data/env_recipes/xxx.json", "相对仓库根路径"),
]
for r, row in enumerate(rows, start=2):
    for c, v in enumerate(row, start=1):
        ws2.cell(row=r, column=c, value=v)
style_range(ws2, 2, len(rows) + 1, 1, len(headers), fill=fill_input, font=font_body, align=left)

set_widths(ws2, {"A": 16, "B": 12, "C": 10, "D": 22, "E": 38, "F": 46})
ws2.freeze_panes = "A2"

# ---------- Sheet 3: 用例模板 ----------
ws3 = wb.create_sheet("用例模板")
ws3.sheet_view.showGridLines = False
ws3["A1"] = "JSON 用例模板"
ws3["A1"].font = font_title
ws3.merge_cells("A1:E1")

json_text = """{
  "id": "case_<module>_<axis>_<short_name>",
  "axis": "target_prediction | state_rolling | counterfactual | generalization",
  "description": "说明这条用例在验证什么",
  "input": {
    "crop": "作物名",
    "zone_id": "分区 id",
    "start_date": "YYYY-MM-DD",
    "adapt_score": 0.85,
    "growth_plan": {},
    "env_recipe": {},
    "env_recipe_path": "data/env_recipes/xxx.json"
  },
  "perturbation": {
    "字段名": {
      "before": "扰动前的值",
      "after": "扰动后的值"
    }
  },
  "expected": {
    "target_prediction": {},
    "state_rolling": {},
    "counterfactual": {},
    "generalization": {}
  },
  "provenance": {
    "note": "为什么需要这条用例"
  }
}"""

ws3["A3"] = "最小 JSON 结构"
ws3["A3"].font = font_section
ws3["A4"] = json_text
ws3["A4"].font = font_mono
ws3["A4"].alignment = left
ws3.merge_cells("A4:E18")
style_range(ws3, 3, 4, 1, 5, fill=fill_section, font=font_section, align=center)
style_range(ws3, 5, 18, 1, 5, fill=PatternFill("solid", fgColor="FF0F172A"), font=Font(name="Consolas", size=10, color="FFE2E8F0"), align=left)

ws3["A20"] = "四轴说明"
ws3["A20"].font = font_section
ws3["A21"] = "target_prediction：验证预测值是否落在容差范围内。"
ws3["A22"] = "state_rolling：验证状态推进的数学性质。"
ws3["A23"] = "counterfactual：验证扰动后的变化方向。"
ws3["A24"] = "generalization：验证跨分区、跨情境的泛化能力。"
style_range(ws3, 20, 24, 1, 1, fill=fill_section, font=font_section, align=center)
style_range(ws3, 21, 24, 2, 5, fill=fill_input, font=font_body, align=left)

set_widths(ws3, {"A": 34, "B": 12, "C": 12, "D": 12, "E": 12})
ws3.row_dimensions[4].height = 220
ws3.row_dimensions[21].height = 22
ws3.row_dimensions[22].height = 22
ws3.row_dimensions[23].height = 22
ws3.row_dimensions[24].height = 22

# ---------- Sheet 4: 四轴判据 ----------
ws4 = wb.create_sheet("四轴判据")
ws4.sheet_view.showGridLines = False
headers4 = ["轴", "验证目标", "关键判据", "典型触发场景", "注意事项"]
for i, h in enumerate(headers4, start=1):
    ws4.cell(row=1, column=i, value=h)
style_range(ws4, 1, 1, 1, len(headers4), fill=fill_header, font=font_header, align=center)

rows4 = [
    ("target_prediction", "预测值是否落在容差范围内", "yield_estimate_g、adjusted_duration_days、harvest_date 及其他模块自定义预测指标", "产量、采收日、关键指标估算", "expected 里按轴填写，不要混入其他轴字段"),
    ("state_rolling", "状态推进的数学性质", "adjusted_duration_days >= 1；harvest_date >= start_date；yield_estimate_g > 0", "滚动累计、单调性、边界合法性", "只验证性质，不做精度评价"),
    ("counterfactual", "扰动后的变化方向", "direction: increase/decrease/change；min_delta_pct", "温度、湿度、日照、EC、基质、养分变化", "perturbation 必填 before / after"),
    ("generalization", "跨分区、跨情境的泛化能力", "all_zones_legal；yields_distinct；metric_bounds", "同一作物跨多个气候区", "不同分区产量不应完全相同"),
]
for r, row in enumerate(rows4, start=2):
    for c, v in enumerate(row, start=1):
        ws4.cell(row=r, column=c, value=v)
style_range(ws4, 2, len(rows4) + 1, 1, len(headers4), fill=fill_input, font=font_body, align=left)
style_range(ws4, 1, len(rows4) + 1, 1, 1, fill=fill_warn, font=font_section, align=center)

set_widths(ws4, {"A": 22, "B": 28, "C": 42, "D": 28, "E": 28})
ws4.freeze_panes = "A2"

# ---------- Sheet 5: 提交清单 ----------
ws5 = wb.create_sheet("提交清单")
ws5.sheet_view.showGridLines = False
headers5 = ["序号", "检查项", "通过标准", "状态"]
for i, h in enumerate(headers5, start=1):
    ws5.cell(row=1, column=i, value=h)
style_range(ws5, 1, 1, 1, len(headers5), fill=fill_header, font=font_header, align=center)

rows5 = [
    (1, "是否已补齐四轴各 1 条用例", "target_prediction / state_rolling / counterfactual / generalization 齐全", ""),
    (2, "是否已写入 cases/ 目录", "真实用例文件位于 benchmarks/agri_recipe_bench/cases/", ""),
    (3, "check_bench_coverage 是否通过", "python scripts/check_bench_coverage.py 通过", ""),
    (4, "run_bench 是否通过", "python benchmarks/agri_recipe_bench/run_bench.py 通过", ""),
    (5, "provenance 是否写明判定动机", "每条用例都可解释‘为什么要验这个’", ""),
]
for r, row in enumerate(rows5, start=2):
    for c, v in enumerate(row, start=1):
        ws5.cell(row=r, column=c, value=v)
style_range(ws5, 2, len(rows5) + 1, 1, len(headers5), fill=fill_input, font=font_body, align=left)

set_widths(ws5, {"A": 8, "B": 34, "C": 44, "D": 14})
ws5.freeze_panes = "A2"

# Save
wb.save(OUT)
print(OUT)
